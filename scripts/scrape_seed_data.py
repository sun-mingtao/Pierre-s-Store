# -*- coding: utf-8 -*-
import argparse
import hashlib
import re
import time
import urllib.parse
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "pierres_seed_data.xlsx"

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
}

SHOP_SOURCES = [
    (
        "灰机Wiki",
        "https://xinglugu.huijiwiki.com/wiki/index.php?title=%E7%9A%AE%E5%9F%83%E5%B0%94%E7%9A%84%E6%9D%82%E8%B4%A7%E5%BA%97&action=raw",
    ),
    (
        "B站中文百科",
        "https://wiki.biligame.com/stardewvalley/index.php?title=%E7%9A%AE%E5%9F%83%E5%B0%94%E7%9A%84%E6%9D%82%E8%B4%A7%E5%BA%97&action=raw",
    ),
]

RAW_PAGE_BASES = [
    "https://xinglugu.huijiwiki.com/wiki/index.php?title={title}&action=raw",
    "https://wiki.biligame.com/stardewvalley/index.php?title={title}&action=raw",
]

SECTIONS = {
    "春季": "==春季出售货物==",
    "夏季": "==夏季出售货物==",
    "秋季": "==秋季出售货物==",
}

SEASON_MAP = {"Spring": "春季", "Summer": "夏季", "Fall": "秋季", "Autumn": "秋季", "Winter": "冬季"}

FALLBACK_DETAILS = {
    "防风草种子": ("防风草", 4, 0, 10, 70, "single"),
    "青豆种子": ("青豆", 10, 3, 30, 80, "multiple"),
    "花椰菜种子": ("花椰菜", 12, 0, 40, 350, "single"),
    "土豆种子": ("土豆", 6, 0, 25, 160, "single"),
    "郁金香球茎": ("郁金香", 6, 0, 10, 60, "single"),
    "甘蓝种子": ("甘蓝", 6, 0, 35, 220, "single"),
    "蓝爵士种子": ("蓝爵士", 7, 0, 15, 100, "single"),
    "大蒜种子": ("大蒜", 4, 0, 20, 120, "single"),
    "稻苗": ("未碾米", 8, 0, 20, 60, "single"),
    "甜瓜种子": ("甜瓜", 12, 0, 40, 500, "single"),
    "西红柿种子": ("西红柿", 11, 4, 25, 120, "multiple"),
    "蓝莓种子": ("蓝莓", 13, 4, 40, 100, "multiple"),
    "辣椒种子": ("辣椒", 5, 3, 20, 80, "multiple"),
    "小麦种子": ("小麦", 4, 0, 5, 50, "single"),
    "萝卜种子": ("萝卜", 6, 0, 20, 180, "single"),
    "虞美人种子": ("虞美人", 7, 0, 50, 280, "single"),
    "夏季亮片种子": ("夏季亮片", 8, 0, 25, 180, "single"),
    "啤酒花种子": ("啤酒花", 11, 1, 30, 50, "multiple"),
    "玉米种子": ("玉米", 14, 4, 75, 100, "multiple"),
    "向日葵种子": ("向日葵", 8, 0, 100, 160, "single"),
    "红叶卷心菜种子": ("红叶卷心菜", 9, 0, 50, 520, "single"),
    "茄子种子": ("茄子", 5, 5, 10, 120, "multiple"),
    "南瓜种子": ("南瓜", 13, 0, 50, 640, "single"),
    "小白菜种子": ("小白菜", 4, 0, 25, 160, "single"),
    "山药种子": ("山药", 10, 0, 30, 320, "single"),
    "蔓越莓种子": ("蔓越莓", 7, 5, 120, 150, "multiple"),
    "玫瑰仙子种子": ("玫瑰仙子", 12, 0, 100, 580, "single"),
    "苋菜种子": ("苋菜", 7, 0, 35, 300, "single"),
    "葡萄种子": ("葡萄", 10, 3, 30, 160, "multiple"),
    "洋蓟种子": ("洋蓟", 8, 0, 15, 320, "single"),
}


def request_text(url):
    res = requests.get(url, headers=HEADERS, timeout=18)
    if res.status_code >= 400 or not res.text.strip():
        raise RuntimeError(f"HTTP {res.status_code}")
    return res.text


def fetch_shop_raw():
    last_error = None
    for source, url in SHOP_SOURCES:
        try:
            text = request_text(url)
            if "春季出售货物" in text:
                return source, text
        except Exception as exc:
            last_error = exc
    raise RuntimeError(f"无法下载商店页面：{last_error}")


def fetch_raw_page(title):
    encoded = urllib.parse.quote(title.replace(" ", "_"))
    for base in RAW_PAGE_BASES:
        url = base.format(title=encoded)
        try:
            text = request_text(url)
        except Exception:
            continue
        redirect = re.search(r"#重定向\s*\[\[([^\]]+)\]\]", text)
        if redirect:
            return fetch_raw_page(redirect.group(1))
        if "DOCTYPE html" not in text[:200].upper():
            return text
        time.sleep(0.15)
    return ""


def section_text(raw, heading):
    if heading not in raw:
        return ""
    part = raw.split(heading, 1)[1]
    return part.split("\n==", 1)[0]


def wiki_link_name(cell):
    links = re.findall(r"\[\[([^\]|]+)(?:\|[^\]]+)?\]\]", cell)
    for link in links:
        if not link.startswith("File:") and not link.startswith("文件:"):
            return link
    return ""


def image_url(file_name):
    if not file_name:
        return ""
    clean = file_name.replace(" ", "_")
    digest = hashlib.md5(clean.encode("utf-8")).hexdigest()
    return f"https://stardewvalleywiki.com/mediawiki/images/{digest[0]}/{digest[:2]}/{urllib.parse.quote(clean)}"


def parse_price(text):
    match = re.search(r"\{\{\s*[Pp]rice\s*\|\s*(\d+)", text)
    return int(match.group(1)) if match else 0


def parse_day(text):
    match = re.search(r"(\d+)\s*天", text or "")
    return int(match.group(1)) if match else 0


def parse_field(raw, field):
    match = re.search(rf"\|\s*{re.escape(field)}\s*=\s*([^\n]+)", raw)
    return match.group(1).strip() if match else ""


def parse_shop_rows(raw):
    rows = []
    for season, heading in SECTIONS.items():
        text = section_text(raw, heading)
        for block in text.split("|-"):
            if "[[File:" not in block or "{{price" not in block.lower():
                continue
            file_match = re.search(r"\[\[File:([^\]|]+)", block)
            name = wiki_link_name(block)
            price = parse_price(block)
            if not name or not price:
                continue
            rows.append(
                {
                    "name": name,
                    "image_file": file_match.group(1) if file_match else "",
                    "image_url": image_url(file_match.group(1) if file_match else ""),
                    "selling_season": season,
                    "seed_buy_price": price,
                }
            )
    return rows


def parse_seed_detail(seed_name, season):
    fallback = FALLBACK_DETAILS.get(seed_name)
    seed_raw = fetch_raw_page(seed_name)
    crop_token = ""
    growth = 0
    seed_sell = 0
    if seed_raw:
        crop_field = parse_field(seed_raw, "crop")
        crop_name_match = re.search(r"\{\{Name\|([^}|]+)", crop_field)
        crop_link_match = re.search(r"\[\[([^\]|]+)", crop_field)
        crop_token = (crop_name_match or crop_link_match).group(1) if (crop_name_match or crop_link_match) else ""
        growth = parse_day(parse_field(seed_raw, "growth"))
        seed_sell = int(re.sub(r"\D", "", parse_field(seed_raw, "sellprice") or "0") or 0)
        raw_season = parse_field(seed_raw, "season")
        season_match = re.search(r"\{\{Season\|([^}|]+)", raw_season)
        if season_match:
            season = SEASON_MAP.get(season_match.group(1), season)

    crop_name, crop_price, regrowth = "", 0, 0
    if crop_token:
        crop_raw = fetch_raw_page(crop_token)
        if crop_raw:
            crop_name = parse_field(crop_raw, "name") or crop_token
            base_price = int(re.sub(r"\D", "", parse_field(crop_raw, "sellprice") or "0") or 0)
            crop_price = base_price * 2 if base_price else 0
            regrowth = parse_day(parse_field(crop_raw, "regrowth"))

    if fallback:
        fallback_crop, fallback_growth, fallback_regrowth, fallback_seed_sell, fallback_crop_price, fallback_type = fallback
        crop_name = crop_name or fallback_crop
        growth = growth or fallback_growth
        regrowth = regrowth or fallback_regrowth
        seed_sell = seed_sell or fallback_seed_sell
        crop_price = crop_price or fallback_crop_price
        harvest_type = "multiple" if regrowth else fallback_type
    else:
        harvest_type = "multiple" if regrowth else "single"

    description = f"{season}种子，{growth} 天成熟"
    if harvest_type == "multiple":
        description += f"，成熟后每 {regrowth} 天再次收获"
    description += f"，可收获{crop_name or '作物'}。"
    return {
        "selling_season": season,
        "growth_time_single": growth,
        "growth_time_multiple": regrowth,
        "seed_sell_price": seed_sell,
        "mature_crop_name": crop_name,
        "mature_crop_price": crop_price,
        "harvest_type": harvest_type,
        "description": description,
    }


def build_rows():
    source, raw = fetch_shop_raw()
    rows = parse_shop_rows(raw)
    enriched = []
    for row in rows:
        detail = parse_seed_detail(row["name"], row["selling_season"])
        row.update(detail)
        row["is_available"] = 1
        row["source"] = source
        enriched.append(row)
        time.sleep(0.1)
    return enriched


def write_xlsx(rows, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "皮埃尔种子数据"
    headers = [
        "name",
        "image_url",
        "description",
        "selling_season",
        "growth_time_single",
        "growth_time_multiple",
        "seed_buy_price",
        "mature_crop_name",
        "mature_crop_price",
        "harvest_type",
        "is_available",
        "seed_sell_price",
        "source",
    ]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header, "") for header in headers])
    fill = PatternFill("solid", fgColor="78A85A")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
    for column in ws.columns:
        width = max(len(str(cell.value or "")) for cell in column) + 2
        ws.column_dimensions[column[0].column_letter].width = min(max(width, 12), 46)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main():
    parser = argparse.ArgumentParser(description="爬取皮埃尔商店种子数据并生成 xlsx。")
    parser.add_argument("--out", default=str(OUT), help="输出 xlsx 路径")
    args = parser.parse_args()
    rows = build_rows()
    write_xlsx(rows, Path(args.out))
    print(f"已生成 {args.out}，共 {len(rows)} 条种子数据。")


if __name__ == "__main__":
    main()

