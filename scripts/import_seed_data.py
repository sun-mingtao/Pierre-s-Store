# -*- coding: utf-8 -*-
import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook
from werkzeug.security import generate_password_hash

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app import app  # noqa: E402
from db import execute, get_connection, select_one  # noqa: E402


DEFAULT_XLSX = ROOT / "data" / "pierres_seed_data.xlsx"


def as_number(value, default=0):
    if value is None or value == "":
        return default
    return float(value)


def import_rows(xlsx_path):
    """
    从 xlsx 文件导入种子数据到数据库。

    根据实体分离原则，种子数据分两张表存储：
    - crops 表：作物生长属性（名称、售价、收获类型、生长时间）
    - seeds 表：种子交易属性（名称、图片、描述、季节、价格、上架状态）
    两表通过 seeds.crop_id 外键关联。
    """
    wb = load_workbook(xlsx_path)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    count = 0
    with app.app_context():
        conn = get_connection()
        for values in ws.iter_rows(min_row=2, values_only=True):
            row = dict(zip(headers, values))
            name = (row.get("name") or "").strip()
            if not name:
                continue

            # ---- 第一步：处理作物（crops）表 ----
            crop_name = (row.get("mature_crop_name") or "").strip()
            crop_price = as_number(row.get("mature_crop_price"))
            harvest_type = row.get("harvest_type") or "single"
            growth_single = int(as_number(row.get("growth_time_single")))
            growth_multiple = int(as_number(row.get("growth_time_multiple")))

            # 查找是否已存在同名作物
            existing_crop = select_one(
                "SELECT crop_id FROM crops WHERE name = %s",
                [crop_name],
            )
            if existing_crop:
                crop_id = existing_crop["crop_id"]
                execute(
                    """
                    UPDATE crops
                    SET sell_price = %s, harvest_type = %s,
                        growth_time_single = %s, growth_time_multiple = %s
                    WHERE crop_id = %s
                    """,
                    [crop_price, harvest_type, growth_single, growth_multiple, crop_id],
                )
            else:
                execute(
                    """
                    INSERT INTO crops (name, sell_price, harvest_type, growth_time_single, growth_time_multiple)
                    VALUES (%s, %s, %s, %s, %s)
                    """,
                    [crop_name, crop_price, harvest_type, growth_single, growth_multiple],
                )
                crop_id = conn.insert_id()

            # ---- 第二步：处理种子（seeds）表 ----
            existing_seed = select_one(
                "SELECT product_id FROM seeds WHERE name = %s",
                [name],
            )
            if existing_seed:
                execute(
                    """
                    UPDATE seeds
                    SET image_url = %s, description = %s, selling_season = %s,
                        seed_buy_price = %s, seed_sell_price = %s,
                        is_available = %s, crop_id = %s
                    WHERE product_id = %s
                    """,
                    [
                        row.get("image_url") or "",
                        row.get("description") or "",
                        row.get("selling_season") or "全年",
                        as_number(row.get("seed_buy_price")),
                        as_number(row.get("seed_sell_price")),
                        int(as_number(row.get("is_available"), 1)),
                        crop_id,
                        existing_seed["product_id"],
                    ],
                )
            else:
                execute(
                    """
                    INSERT INTO seeds
                    (name, image_url, description, selling_season,
                     seed_buy_price, seed_sell_price, is_available, crop_id)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
                    """,
                    [
                        name,
                        row.get("image_url") or "",
                        row.get("description") or "",
                        row.get("selling_season") or "全年",
                        as_number(row.get("seed_buy_price")),
                        as_number(row.get("seed_sell_price")),
                        int(as_number(row.get("is_available"), 1)),
                        crop_id,
                    ],
                )
            count += 1
    return count


def create_demo_users():
    demo = [
        ("皮埃尔", "pierre123", "owner", 99999),
        ("农场主", "player123", "player", 1500),
    ]
    created = 0
    with app.app_context():
        for username, password, role, coins in demo:
            existing = select_one("SELECT user_id FROM users WHERE username = %s", [username], ["user_id"])
            if existing:
                continue
            execute(
                "INSERT INTO users (username, password, role, coins) VALUES (%s, %s, %s, %s)",
                [username, generate_password_hash(password), role, coins],
            )
            created += 1
    return created


def main():
    parser = argparse.ArgumentParser(description="将 xlsx 种子数据导入 MySQL。")
    parser.add_argument("--xlsx", default=str(DEFAULT_XLSX), help="xlsx 文件路径")
    parser.add_argument("--demo-users", action="store_true", help="创建演示账号")
    args = parser.parse_args()
    count = import_rows(Path(args.xlsx))
    created = create_demo_users() if args.demo_users else 0
    print(f"已导入/更新 {count} 条种子数据。")
    if args.demo_users:
        print(f"已创建 {created} 个演示账号。")


if __name__ == "__main__":
    main()

